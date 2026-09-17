from __future__ import annotations

from collections import Counter
from io import BytesIO
from pathlib import Path
import json
import re
from zipfile import BadZipFile, ZipFile

import pymupdf
from docx import Document
from lxml import etree
from pydantic import BaseModel, Field

from resume_mvp.config import settings
from resume_mvp.domain import (
    CustomSection,
    EducationEntry,
    Fact,
    LayoutProfile,
    ProjectEntry,
    ResumeDocument,
    SkillGroup,
    SourcedText,
    WorkExperienceEntry,
)
from resume_mvp.profile import facts_from_resume
from resume_mvp.text_tidy import tidy_paste_artifacts


class ResumeImportError(ValueError):
    pass


class ImportResult(BaseModel):
    resume: ResumeDocument
    facts: list[Fact]
    layout_profile: LayoutProfile
    quality_score: float = Field(ge=0, le=1)
    warnings: list[str] = Field(default_factory=list)


_MIME_BY_EXTENSION = {
    ".json": {"application/json", "text/json", "text/plain", "application/octet-stream"},
    ".txt": {"text/plain", "application/octet-stream"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
    },
    ".pdf": {"application/pdf", "application/octet-stream"},
}

_SECTION_HEADINGS = {
    "工作经历": "work",
    "工作经验": "work",
    "职业经历": "work",
    "实习经历": "work",
    "实习经验": "work",
    "项目经历": "project",
    "项目经验": "project",
    "教育经历": "education",
    "教育背景": "education",
    "专业技能": "skills",
    "技能": "skills",
    "证书": "certificate",
    "资格证书": "certificate",
    "奖项": "certificate",
    "荣誉奖项": "certificate",
}


def import_resume(filename: str, content_type: str, data: bytes) -> ImportResult:
    extension = Path(filename).suffix.lower()
    if extension not in _MIME_BY_EXTENSION:
        raise ResumeImportError("仅支持 JSON、PDF、DOCX 和 TXT 文件")
    if content_type not in _MIME_BY_EXTENSION[extension]:
        raise ResumeImportError("文件类型不匹配，请检查扩展名")
    if not data:
        raise ResumeImportError("文件内容为空")
    if len(data) > settings.max_upload_bytes:
        raise ResumeImportError("文件超过 10 MiB 限制")

    if extension == ".json":
        return _extract_json(data)
    extraction_warnings: list[str] = []
    if extension == ".txt":
        lines, layout = _extract_txt(data)
    elif extension == ".docx":
        lines, layout, extraction_warnings = _extract_docx(data)
    else:
        lines, layout, extraction_warnings = _extract_pdf(data)

    lines = [
        cleaned
        for line in lines
        if line and (cleaned := tidy_paste_artifacts(line).strip())
    ]
    if not lines:
        raise ResumeImportError("未提取到可用文字；扫描版 PDF 暂不支持 OCR")

    resume, facts, recognized_sections = _structure_lines(lines, layout)
    quality = _quality_score(lines, resume, recognized_sections)
    warnings: list[str] = list(extraction_warnings)
    if extension == ".pdf":
        warnings.append("PDF 模板将按文字块和样式特征近似重建")
    if recognized_sections == 0 and len(lines) > 3:
        warnings.append("未识别到标准章节，请确认结构化结果")
    if quality < 0.55:
        warnings.append("解析质量较低，保存前请逐项确认")

    return ImportResult(
        resume=resume,
        facts=facts,
        layout_profile=layout,
        quality_score=quality,
        warnings=warnings,
    )


def _extract_json(data: bytes) -> ImportResult:
    """Restore either an exported profile envelope or a bare ResumeDocument."""
    try:
        payload = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise ResumeImportError("JSON 文件无效或不是 UTF-8 编码") from error
    if not isinstance(payload, dict):
        raise ResumeImportError("JSON 顶层必须是对象")
    raw_resume = payload.get("resume", payload)
    try:
        resume = ResumeDocument.model_validate(raw_resume)
    except Exception as error:
        raise ResumeImportError("JSON 不符合经历库结构") from error
    facts = facts_from_resume(resume)
    return ImportResult(
        resume=resume,
        facts=facts,
        layout_profile=resume.layout_profile,
        quality_score=1.0,
        warnings=["已从 JSON 备份恢复结构化经历；保存前请核对内容"],
    )


def _extract_txt(data: bytes) -> tuple[list[str], LayoutProfile]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ResumeImportError("TXT 文件必须使用 UTF-8 编码") from error
    return text.splitlines(), LayoutProfile(source_kind="txt")


def _extract_docx(data: bytes) -> tuple[list[str], LayoutProfile, list[str]]:
    try:
        document = Document(BytesIO(data))
    except Exception as error:
        raise ResumeImportError("DOCX 文件无法读取") from error

    paragraph_lines: list[str] = []
    fonts: Counter[str] = Counter()
    sizes: list[float] = []
    colors: Counter[str] = Counter()
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraph_lines.append(paragraph.text.strip())
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            if run.font.name:
                fonts[run.font.name] += len(run.text)
            if run.font.size:
                sizes.append(run.font.size.pt)
            if run.font.color and run.font.color.rgb:
                colors[f"#{str(run.font.color.rgb).upper()}"] += len(run.text)

    xml_lines = _extract_docx_xml_lines(data)
    lines = xml_lines or paragraph_lines
    warnings = (
        ["已读取 DOCX 正文、表格、文本框及页眉页脚；请核对复杂浮动布局的顺序"]
        if xml_lines
        else ["DOCX 复杂结构读取失败，已降级为普通段落提取"]
    )
    font = fonts.most_common(1)[0][0] if fonts else ""
    accent = _preferred_color(colors)
    return lines, LayoutProfile(
        source_kind="docx",
        font_family=font,
        heading_font_family=font,
        accent_color=accent,
        base_font_size=min(sizes) if sizes else None,
        imported=True,
    ), warnings


def _extract_docx_xml_lines(data: bytes) -> list[str]:
    """Read paragraphs from body/tables/textboxes/header/footer in document order."""
    word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    mc_ns = "http://schemas.openxmlformats.org/markup-compatibility/2006"

    def inside_fallback(node: etree._Element) -> bool:
        parent = node.getparent()
        while parent is not None:
            qname = etree.QName(parent)
            if qname.namespace == mc_ns and qname.localname == "Fallback":
                return True
            parent = parent.getparent()
        return False

    def paragraph_text(paragraph: etree._Element) -> str:
        chunks: list[str] = []
        for node in paragraph.iter():
            qname = etree.QName(node)
            if qname.namespace != word_ns or inside_fallback(node):
                continue
            if qname.localname == "t" and node.text:
                chunks.append(node.text)
            elif qname.localname == "tab":
                chunks.append("\t")
            elif qname.localname in {"br", "cr"}:
                chunks.append("\n")
        return "".join(chunks).strip()

    def walk(node: etree._Element, output: list[str]) -> None:
        for child in node:
            qname = etree.QName(child)
            if qname.namespace == mc_ns and qname.localname == "AlternateContent":
                choice = next(
                    (
                        candidate
                        for candidate in child
                        if etree.QName(candidate).namespace == mc_ns
                        and etree.QName(candidate).localname == "Choice"
                    ),
                    None,
                )
                fallback = next(
                    (
                        candidate
                        for candidate in child
                        if etree.QName(candidate).namespace == mc_ns
                        and etree.QName(candidate).localname == "Fallback"
                    ),
                    None,
                )
                selected = choice if choice is not None else fallback
                if selected is not None:
                    walk(selected, output)
                continue
            if qname.namespace == word_ns and qname.localname == "p":
                text = paragraph_text(child)
                if text:
                    output.extend(part for part in text.splitlines() if part.strip())
                continue
            walk(child, output)

    try:
        with ZipFile(BytesIO(data)) as archive:
            names = [
                name
                for name in archive.namelist()
                if name == "word/document.xml"
                or (
                    name.startswith(("word/header", "word/footer"))
                    and name.endswith(".xml")
                )
            ]
            result: list[str] = []
            for name in sorted(names, key=lambda item: (item != "word/document.xml", item)):
                try:
                    root = etree.fromstring(archive.read(name))
                except (KeyError, etree.XMLSyntaxError):
                    continue
                walk(root, result)
            return result
    except (BadZipFile, OSError):
        return []


def _extract_pdf(data: bytes) -> tuple[list[str], LayoutProfile, list[str]]:
    try:
        document = pymupdf.open(stream=data, filetype="pdf")
    except Exception as error:
        raise ResumeImportError("PDF 文件无法读取") from error

    primary_lines: list[str] = []
    fallback_lines: list[str] = []
    fonts: Counter[str] = Counter()
    sizes: list[float] = []
    colors: Counter[str] = Counter()
    x_positions: list[float] = []
    warnings: list[str] = []
    page_limit = min(document.page_count, 50)
    if document.page_count > page_limit:
        warnings.append(f"PDF 共 {document.page_count} 页，仅解析前 {page_limit} 页")
    for page_index in range(page_limit):
        page = document[page_index]
        payload = page.get_text("dict")
        for block in payload.get("blocks", []):
            if "lines" not in block:
                continue
            x_positions.append(float(block.get("bbox", [0])[0]))
            for line in block["lines"]:
                text = "".join(span.get("text", "") for span in line["spans"]).strip()
                if text:
                    primary_lines.append(text)
                for span in line["spans"]:
                    value = span.get("text", "")
                    if not value.strip():
                        continue
                    fonts[str(span.get("font", ""))] += len(value)
                    sizes.append(float(span.get("size", 0)))
                    color = int(span.get("color", 0)) & 0xFFFFFF
                    if color:
                        colors[f"#{color:06X}"] += len(value)
        blocks = [
            block
            for block in page.get_text("blocks")
            if len(block) >= 5 and str(block[4]).strip()
        ]
        if blocks:
            xs = [float(block[0]) for block in blocks]
            two_column = _looks_two_column(xs)
            ordered = sorted(
                blocks,
                key=(
                    (lambda block: (0 if float(block[0]) < page.rect.width * 0.45 else 1, float(block[1]), float(block[0])))
                    if two_column
                    else (lambda block: (float(block[1]), float(block[0])))
                ),
            )
            for block in ordered:
                fallback_lines.extend(
                    line.strip()
                    for line in str(block[4]).splitlines()
                    if line.strip()
                )
    document.close()

    lines = primary_lines
    if _line_extraction_score(fallback_lines) > _line_extraction_score(primary_lines) * 1.15:
        lines = fallback_lines
        warnings.append("PDF 主读取顺序质量较低，已自动使用分块排序结果")
    collapsed = _collapse_fragmented_lines(lines)
    if collapsed != lines:
        lines = collapsed
        warnings.append("PDF 存在字符级断行，已合并连续碎片")

    columns = 2 if _looks_two_column(x_positions) else 1
    return lines, LayoutProfile(
        source_kind="pdf",
        font_family=fonts.most_common(1)[0][0] if fonts else "",
        accent_color=_preferred_color(colors),
        base_font_size=min((size for size in sizes if size > 0), default=None),
        columns=columns,
        imported=True,
    ), warnings


def _line_extraction_score(lines: list[str]) -> float:
    if not lines:
        return 0.0
    chars = sum(len(line.strip()) for line in lines)
    fragments = sum(len(line.strip()) <= 2 for line in lines)
    return max(0.0, chars - fragments * 4)


def _collapse_fragmented_lines(lines: list[str]) -> list[str]:
    if len(lines) < 6:
        return lines
    short_ratio = sum(len(line.strip()) <= 2 for line in lines) / len(lines)
    if short_ratio < 0.3:
        return lines
    result: list[str] = []
    fragments: list[str] = []

    def flush() -> None:
        if fragments:
            result.append("".join(fragments))
            fragments.clear()

    for line in lines:
        stripped = line.strip()
        if 0 < len(stripped) <= 2:
            fragments.append(stripped)
            continue
        flush()
        if stripped:
            result.append(stripped)
    flush()
    return result


def _preferred_color(colors: Counter[str]) -> str:
    for color, _ in colors.most_common():
        if color not in {"#000000", "#FFFFFF"}:
            return color
    return ""


def _looks_two_column(x_positions: list[float]) -> bool:
    if len(x_positions) < 6:
        return False
    left = [x for x in x_positions if x < 220]
    right = [x for x in x_positions if x > 260]
    return len(left) >= 2 and len(right) >= 2


def _structure_lines(
    lines: list[str], layout: LayoutProfile
) -> tuple[ResumeDocument, list[Fact], int]:
    resume = ResumeDocument(layout_profile=layout)
    facts = [
        Fact(
            category="简历原文",
            statement=line,
            source_type="upload",
            source_location=f"第 {index + 1} 行",
        )
        for index, line in enumerate(lines)
    ]
    fact_by_line = {fact.statement: fact for fact in facts}

    resume.basics.name = lines[0] if len(lines[0]) <= 16 else ""
    if len(lines) > 1 and lines[1] not in _SECTION_HEADINGS:
        resume.basics.target_role = _sourced(lines[1], fact_by_line)

    section: str | None = None
    recognized = 0
    raw: list[SourcedText] = []
    work: WorkExperienceEntry | None = None
    project: ProjectEntry | None = None
    education: EducationEntry | None = None

    for index, line in enumerate(lines):
        if index < 2:
            continue
        normalized = re.sub(r"[：:\s]", "", line)
        heading = _SECTION_HEADINGS.get(normalized)
        if heading:
            section = heading
            recognized += 1
            work = None
            project = None
            education = None
            continue

        sourced = _sourced(line, fact_by_line)
        if section == "work":
            if work is None:
                company, title = _split_title(line)
                work = WorkExperienceEntry(company=company, title=title)
                resume.work_experience.append(work)
            else:
                work.bullets.append(sourced)
        elif section == "project":
            if project is None:
                name, role = _split_title(line)
                project = ProjectEntry(name=name, role=role)
                resume.projects.append(project)
            else:
                project.bullets.append(sourced)
        elif section == "education":
            if education is None:
                institution, degree = _split_title(line)
                education = EducationEntry(institution=institution, degree=degree)
                resume.education.append(education)
            else:
                education.highlights.append(sourced)
        elif section == "skills":
            items = [item.strip() for item in re.split(r"[、,，/|｜]", line) if item.strip()]
            resume.skills.append(
                SkillGroup(name="技能", items=[_sourced(item, fact_by_line, fallback=fact_by_line[line]) for item in items])
            )
        else:
            raw.append(sourced)

    if raw:
        resume.custom_sections.append(CustomSection(title="原始内容", items=raw))
        if "custom_sections" not in resume.section_order:
            resume.section_order.append("custom_sections")
    return resume, facts, recognized


def _sourced(
    text: str,
    fact_by_line: dict[str, Fact],
    *,
    fallback: Fact | None = None,
) -> SourcedText:
    fact = fact_by_line.get(text, fallback)
    return SourcedText(
        value=text,
        source_fact_ids=[fact.id] if fact else [],
        origin="upload",
        confidence=0.85,
    )


def _split_title(line: str) -> tuple[str, str]:
    parts = [part.strip() for part in re.split(r"[|｜]", line, maxsplit=1)]
    return (parts[0], parts[1] if len(parts) > 1 else "")


def _quality_score(lines: list[str], resume: ResumeDocument, recognized: int) -> float:
    score = 0.25
    score += min(len(lines) / 20, 0.25)
    score += min(recognized * 0.1, 0.3)
    if resume.basics.name:
        score += 0.1
    if resume.basics.target_role.value:
        score += 0.1
    return round(min(score, 1.0), 2)
