from __future__ import annotations

from base64 import b64decode
from io import BytesIO
import json
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from resume_mvp.domain import (
    ApplicationType,
    Fact,
    JobAnalysis,
    JobProject,
    ResumeDocument,
    ResumeVersion,
    utc_now,
)
from resume_mvp.layout_tidy import skill_lines_for_export, tidy_resume_for_layout


_TEMPLATE_STYLES = {
    "overleaf-cn": {"font": "Helvetica Neue", "accent": "003E74", "margin": 1.2, "classic": True},
    "classic-cn": {"font": "SimSun", "accent": "111111", "margin": 1.6, "classic": True},
    "clear-single": {"font": "Microsoft YaHei", "accent": "2457D6", "margin": 1.7},
    "pro-double": {"font": "Microsoft YaHei", "accent": "214A72", "margin": 1.5},
    "project-focus": {"font": "Microsoft YaHei", "accent": "7A3E8E", "margin": 1.6},
    "career-depth": {"font": "SimSun", "accent": "8A4C2A", "margin": 1.8},
}


def build_docx(
    resume: ResumeDocument,
    template_id: str,
    application_type: ApplicationType = "experienced",
) -> bytes:
    style = dict(_TEMPLATE_STYLES.get(template_id, _TEMPLATE_STYLES["classic-cn"]))
    classic = bool(style.get("classic"))
    body_size = 10.5 if classic else 10.0
    if resume.layout_profile.imported and template_id in {"clear-single", "classic-cn"}:
        imported_font = resume.layout_profile.font_family.strip()
        imported_accent = resume.layout_profile.accent_color.strip()
        if imported_font and len(imported_font) <= 100:
            style["font"] = imported_font
        if re.fullmatch(r"#[0-9A-Fa-f]{6}", imported_accent) and not classic:
            style["accent"] = imported_accent[1:].upper()
        if resume.layout_profile.base_font_size is not None:
            body_size = min(max(resume.layout_profile.base_font_size, 8), 12)
    compact = application_type in {"campus", "internship"} or classic
    resume = tidy_resume_for_layout(resume)
    document = Document()
    section = document.sections[0]
    if classic:
        margin = Cm(style["margin"])
    elif compact:
        margin = Cm(min(style["margin"], 1.25))
    else:
        margin = Cm(style["margin"])
    section.top_margin = section.bottom_margin = margin
    section.left_margin = section.right_margin = margin

    normal = document.styles["Normal"]
    normal.font.name = style["font"]
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), style["font"])
    normal.font.size = Pt(10.5 if classic else (9 if compact else body_size))
    normal.paragraph_format.space_after = Pt(1 if compact else 4)
    normal.paragraph_format.line_spacing = 1.05 if classic else (1.0 if compact else 1.15)

    _write_header(document, resume, style, classic=classic)
    _write_basic_info_block(document, resume, style, compact=compact, classic=classic)

    # Campus/internship paper resumes usually skip a free-text summary block.
    show_summary = application_type == "experienced" and bool(resume.basics.summary.value.strip())
    if show_summary:
        _heading(document, "个人简介", style, compact=compact, classic=classic)
        document.add_paragraph(resume.basics.summary.value)

    section_order = list(resume.section_order)
    if template_id == "project-focus" and "projects" in section_order:
        section_order.remove("projects")
        section_order.insert(1, "projects")

    for section_name in section_order:
        if section_name == "work_experience" and resume.work_experience:
            _heading(
                document,
                "职业经历" if classic else "实习工作经历",
                style,
                compact=compact,
                classic=classic,
            )
            for item in resume.work_experience:
                _entry_title(
                    document,
                    item.company,
                    item.title,
                    item.start_date,
                    item.end_date,
                    style,
                    classic=classic,
                )
                _bullets(document, [bullet.value for bullet in item.bullets], classic=classic)
        elif section_name == "projects" and resume.projects:
            _heading(document, "项目经历", style, compact=compact, classic=classic)
            for item in resume.projects:
                _entry_title(
                    document,
                    item.name,
                    item.role,
                    item.start_date,
                    item.end_date,
                    style,
                    classic=classic,
                )
                _bullets(document, [bullet.value for bullet in item.bullets], classic=classic)
        elif section_name == "education" and resume.education:
            _heading(
                document,
                "教育背景" if classic else "教育经历",
                style,
                compact=compact,
                classic=classic,
            )
            for item in resume.education:
                if classic:
                    _classic_education_entry(
                        document,
                        item.institution,
                        item.field,
                        item.degree,
                        item.start_date,
                        item.end_date,
                        style,
                    )
                else:
                    detail = " · ".join(value for value in [item.degree, item.field] if value)
                    _entry_title(
                        document,
                        item.institution,
                        detail,
                        item.start_date,
                        item.end_date,
                        style,
                        classic=False,
                    )
                _bullets(document, [highlight.value for highlight in item.highlights], classic=classic)
        elif section_name == "skills" and resume.skills:
            lines = skill_lines_for_export(resume.skills)
            if lines:
                _heading(document, "专业技能", style, compact=compact, classic=classic)
                _bullets(document, lines, classic=classic)
        elif section_name == "custom_sections":
            for custom in resume.custom_sections:
                _heading(document, custom.title, style, compact=compact, classic=classic)
                _bullets(document, [item.value for item in custom.items], classic=classic)

    if resume.certificates:
        _heading(document, "证书", style, compact=compact, classic=classic)
        _bullets(document, [entry.name for entry in resume.certificates], classic=classic)
    if resume.awards:
        _heading(document, "奖项", style, compact=compact, classic=classic)
        _bullets(document, [entry.name for entry in resume.awards], classic=classic)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def build_resume_json(project: JobProject, version: ResumeVersion) -> bytes:
    payload = {
        "format": "resume-evidence-workbench/v1",
        "exported_at": utc_now().isoformat(),
        "project": {
            "id": project.id,
            "title": project.title,
            "company_name": project.company_name,
            "application_type": project.application_type,
            "job_description": project.job_description,
            "selected_template_id": project.selected_template_id,
        },
        "version_id": version.id,
        "resume": version.resume.model_dump(mode="json"),
        "facts": [fact.model_dump(mode="json") for fact in version.facts],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def build_codex_handoff(
    project: JobProject,
    analysis: JobAnalysis | None,
    resume: ResumeDocument,
    facts: list[Fact],
) -> str:
    safe_resume = resume.model_dump(mode="json")
    safe_resume["basics"]["email"] = ""
    safe_resume["basics"]["phone"] = ""
    safe_resume["basics"]["wechat"] = ""
    if safe_resume["basics"].get("photo_data_url"):
        safe_resume["basics"]["photo_data_url"] = "[已上传证件照]"
    confirmed = [fact.model_dump(mode="json") for fact in facts if fact.user_confirmed]
    analysis_payload = analysis.model_dump(mode="json") if analysis else {}
    return "\n\n".join(
        [
            "# Codex 求职上下文",
            "## 任务\n基于以下 JD 与事实帮助优化中文简历或进行求职训练。",
            f"## 目标\n项目：{project.title}\n公司：{project.company_name or '未填写'}",
            "## JD 摘要\n```json\n"
            + json.dumps(analysis_payload, ensure_ascii=False, indent=2)
            + "\n```",
            "## 已确认事实\n```json\n"
            + json.dumps(confirmed, ensure_ascii=False, indent=2)
            + "\n```",
            "## 当前简历\n```json\n"
            + json.dumps(safe_resume, ensure_ascii=False, indent=2)
            + "\n```",
            "## 待解决问题\n请指出证据不足的岗位要求，并一次提出一个问题。",
            "## 输出约束\n仅使用中文；不得虚构事实；每项改写必须注明对应事实；先给建议，不直接覆盖简历。",
        ]
    )


def _write_header(
    document: Document,
    resume: ResumeDocument,
    style: dict,
    *,
    classic: bool = False,
) -> None:
    if classic:
        _write_classic_header(document, resume, style)
        return

    photo = _decode_photo(resume.basics.photo_data_url)
    if photo is None:
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(4)
        name_run = title.add_run(resume.basics.name or "姓名")
        _style_run(name_run, style["font"], 20, style["accent"], bold=True)
        if resume.basics.target_role.value:
            role = document.add_paragraph()
            role.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _style_run(role.add_run(resume.basics.target_role.value), style["font"], 10, "687386")
        return

    table = document.add_table(rows=1, cols=2)
    table.autofit = True
    left, right = table.rows[0].cells
    name = left.paragraphs[0]
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _style_run(name.add_run(resume.basics.name or "姓名"), style["font"], 20, style["accent"], bold=True)
    if resume.basics.target_role.value:
        role = left.add_paragraph()
        _style_run(role.add_run(resume.basics.target_role.value), style["font"], 10, "687386")
    for line in _basic_info_lines(resume.basics, classic=False):
        info = left.add_paragraph()
        run = info.add_run(line)
        _style_run(run, style["font"], 9, "4B5563")
    photo_paragraph = right.paragraphs[0]
    photo_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = photo_paragraph.add_run()
    _add_cropped_picture(run, photo, width=Cm(2.6), height=Cm(3.4))


def _add_cropped_picture(run, photo: bytes, *, width, height) -> None:
    shape = run.add_picture(BytesIO(photo), width=width, height=height)
    try:
        image = Image.from_blob(photo)
        source_aspect = image.px_width / image.px_height
        target_aspect = width / height
    except (ValueError, ZeroDivisionError):
        return

    crop = {"l": 0, "r": 0, "t": 0, "b": 0}
    if source_aspect > target_aspect:
        horizontal = round((1 - target_aspect / source_aspect) * 50_000)
        crop["l"] = crop["r"] = horizontal
    elif source_aspect < target_aspect:
        vertical = round((1 - source_aspect / target_aspect) * 50_000)
        crop["t"] = crop["b"] = vertical
    else:
        return

    source_rectangle = OxmlElement("a:srcRect")
    for side, value in crop.items():
        if value:
            source_rectangle.set(side, str(value))
    blip_fill = shape._inline.graphic.graphicData.pic.blipFill
    blip_fill.insert(1, source_rectangle)


def _write_classic_header(document: Document, resume: ResumeDocument, style: dict) -> None:
    """Match the Overleaf classic Chinese single-column macros: name / contact / otherInfo / photo."""
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    _style_run(
        title.add_run(resume.basics.name or "姓名"),
        style["font"],
        22,
        style["accent"],
        bold=True,
    )

    for line in _classic_info_lines(resume.basics):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(1)
        _style_run(paragraph.add_run(line), style["font"], 9.5, "333333")

    photo = _decode_photo(resume.basics.photo_data_url)
    if photo is None:
        return
    photo_paragraph = document.add_paragraph()
    photo_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    photo_paragraph.paragraph_format.space_before = Pt(0)
    photo_paragraph.paragraph_format.space_after = Pt(2)
    run = photo_paragraph.add_run()
    page_width = document.sections[0].page_width
    run.add_picture(BytesIO(photo), width=int(page_width * 0.15))


def _write_basic_info_block(
    document: Document,
    resume: ResumeDocument,
    style: dict,
    *,
    compact: bool,
    classic: bool = False,
) -> None:
    if classic:
        # Classic contact/otherInfo already written under the centered name.
        return
    # When a photo header already embeds basics beside the portrait, skip the duplicate block.
    if _decode_photo(resume.basics.photo_data_url) is not None:
        return
    lines = _basic_info_lines(resume.basics, classic=False)
    if not lines:
        return
    for line in lines:
        paragraph = document.add_paragraph(line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(1 if compact else 2)
        for run in paragraph.runs:
            _style_run(run, style["font"], 9 if compact else 9.5, "4B5563")


def _classic_info_lines(basics) -> list[str]:
    """Build LaTeX-like \\contactInfo then \\otherInfo rows (max 4 items each)."""
    contact = " · ".join(
        value
        for value in [
            f"手机：{basics.phone}" if basics.phone.strip() else "",
            f"邮箱：{basics.email}" if basics.email.strip() else "",
            f"微信：{basics.wechat}" if basics.wechat.strip() else "",
        ]
        if value
    )
    other_items = [
        value
        for value in [
            f"性别：{basics.gender}" if basics.gender.strip() else "",
            f"现居：{basics.location}" if basics.location.strip() else "",
            f"政治面貌：{basics.political_status}" if basics.political_status.strip() else "",
            f"生日：{basics.birthday}" if basics.birthday.strip() else "",
        ]
        if value
    ]
    lines: list[str] = []
    if contact:
        lines.append(contact)
    for index in range(0, len(other_items), 4):
        lines.append(" · ".join(other_items[index : index + 4]))
    return lines


def _basic_info_lines(basics, *, classic: bool = False) -> list[str]:
    if classic:
        return _classic_info_lines(basics)
    identity = " · ".join(
        value
        for value in [
            f"性别：{basics.gender}" if basics.gender.strip() else "",
            f"生日：{basics.birthday}" if basics.birthday.strip() else "",
            f"政治面貌：{basics.political_status}" if basics.political_status.strip() else "",
            f"现居：{basics.location}" if basics.location.strip() else "",
        ]
        if value
    )
    contact = " · ".join(
        value
        for value in [
            f"电话：{basics.phone}" if basics.phone.strip() else "",
            f"邮箱：{basics.email}" if basics.email.strip() else "",
            f"微信：{basics.wechat}" if basics.wechat.strip() else "",
        ]
        if value
    )
    return [line for line in [identity, contact] if line]


def _decode_photo(data_url: str) -> bytes | None:
    raw = (data_url or "").strip()
    if not raw.startswith("data:image/") or ";base64," not in raw:
        return None
    try:
        encoded = raw.split(";base64,", 1)[1]
        payload = b64decode(encoded, validate=False)
    except Exception:
        return None
    if len(payload) < 32 or len(payload) > 900_000:
        return None
    return payload


def _heading(
    document: Document,
    text: str,
    style: dict,
    *,
    compact: bool = False,
    classic: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8 if classic else (6 if compact else 10))
    paragraph.paragraph_format.space_after = Pt(3 if classic else (2 if compact else 4))
    _style_run(
        paragraph.add_run(text),
        style["font"],
        12 if classic else (10.5 if compact else 12),
        "111111" if classic else style["accent"],
        bold=True,
    )
    if classic:
        _set_paragraph_bottom_border(paragraph, color="111111")


def _classic_education_entry(
    document: Document,
    institution: str,
    field: str,
    degree: str,
    start: str,
    end: str,
    style: dict,
) -> None:
    """\\datedsubsection{\\textbf{学校}，专业，\\textit{学位}}{dates}"""
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(1)
    usable = (
        document.sections[0].page_width
        - document.sections[0].left_margin
        - document.sections[0].right_margin
    )
    paragraph.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)

    _style_run(paragraph.add_run(institution.strip() or " "), style["font"], 10.5, "111111", bold=True)
    if field.strip():
        _style_run(paragraph.add_run(f"，{field.strip()}"), style["font"], 10.5, "111111")
    if degree.strip():
        _style_run(
            paragraph.add_run(f"，{degree.strip()}"),
            style["font"],
            10.5,
            "111111",
            italic=True,
        )
    dates = " - ".join(value for value in [start, end] if value)
    if dates:
        paragraph.add_run("\t")
        _style_run(paragraph.add_run(dates), style["font"], 10.5, "333333")


def _entry_title(
    document: Document,
    primary: str,
    secondary: str,
    start: str,
    end: str,
    style: dict,
    *,
    classic: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2 if classic else 0)
    paragraph.paragraph_format.space_after = Pt(1 if classic else 0)
    left = primary.strip()
    if secondary.strip():
        left = f"{left}，{secondary.strip()}" if classic else f"{left}  {secondary.strip()}"
    dates = " - ".join(value for value in [start, end] if value) if classic else " – ".join(
        value for value in [start, end] if value
    )

    if classic:
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER

        # Right-align dates on the same line (classic Chinese resume).
        usable = (
            document.sections[0].page_width
            - document.sections[0].left_margin
            - document.sections[0].right_margin
        )
        paragraph.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)
        _style_run(paragraph.add_run(left or " "), style["font"], 10.5, "111111", bold=True)
        if dates:
            paragraph.add_run("\t")
            _style_run(paragraph.add_run(dates), style["font"], 10.5, "333333", bold=False)
        return

    _style_run(paragraph.add_run(primary), style["font"], 10.5, "172033", bold=True)
    if secondary:
        paragraph.add_run(f"  {secondary}")
    if dates:
        paragraph.add_run(f"    {dates}")


def _set_paragraph_bottom_border(paragraph, *, color: str = "111111") -> None:
    from docx.oxml import OxmlElement

    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _bullets(document: Document, values: list[str], *, classic: bool = False) -> None:
    for value in values:
        if not value:
            continue
        paragraph = document.add_paragraph(value, style="List Bullet")
        if classic:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)


def _style_run(
    run,
    font: str,
    size: float,
    color: str,
    *,
    bold: bool = False,
    italic: bool = False,
) -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
