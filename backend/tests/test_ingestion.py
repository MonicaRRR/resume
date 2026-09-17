from io import BytesIO
from zipfile import ZipFile

import fitz
import pytest
from docx import Document
from docx.shared import Pt, RGBColor

from resume_mvp.ingestion import (
    ResumeImportError,
    _collapse_fragmented_lines,
    _extract_docx_xml_lines,
    import_resume,
)


def make_docx() -> bytes:
    document = Document()
    name = document.add_paragraph()
    run = name.add_run("张宁")
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(36, 87, 214)
    document.add_paragraph("后端工程师")
    document.add_paragraph("工作经历")
    document.add_paragraph("示例科技｜后端工程师")
    document.add_paragraph("使用 Python 开发 API")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_blank_pdf() -> bytes:
    document = fitz.open()
    document.new_page()
    return document.tobytes()


def make_structured_docx() -> bytes:
    document = Document()
    document.add_paragraph("李明")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目经历"
    table.cell(0, 1).text = "知识库检索平台"
    table.cell(1, 0).text = "专业技能"
    table.cell(1, 1).text = "Python、Pytest"
    document.sections[0].header.paragraphs[0].text = "liming@example.com"
    document.sections[0].footer.paragraphs[0].text = "作品集"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_txt_import_preserves_lines_as_confirmable_facts() -> None:
    """Catches importers that drop source lines and leave AI without evidence."""
    result = import_resume(
        "resume.txt",
        "text/plain",
        "张宁\n产品经理\n负责增长实验".encode(),
    )

    assert result.resume.basics.name == "张宁"
    assert result.resume.basics.target_role.value == "产品经理"
    assert any("增长实验" in fact.statement for fact in result.facts)
    assert result.quality_score > 0


def test_docx_import_preserves_a_style_profile() -> None:
    """Catches DOCX parsing that extracts text but loses the reusable visual profile."""
    result = import_resume(
        "resume.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        make_docx(),
    )

    assert result.resume.basics.name == "张宁"
    assert result.layout_profile.source_kind == "docx"
    assert result.layout_profile.imported is True
    assert result.layout_profile.font_family == "Microsoft YaHei"
    assert result.layout_profile.accent_color == "#2457D6"


def test_docx_import_reads_tables_headers_and_footers() -> None:
    result = import_resume(
        "structured.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        make_structured_docx(),
    )

    statements = "\n".join(fact.statement for fact in result.facts)
    assert "知识库检索平台" in statements
    assert "Python" in statements
    assert "liming@example.com" in statements
    assert any("表格" in warning for warning in result.warnings)


def test_docx_ooxml_uses_choice_and_skips_duplicate_fallback_textbox() -> None:
    xml = """<?xml version="1.0" encoding="UTF-8"?>
    <w:document
      xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
      xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">
      <w:body>
        <w:p><w:r><w:t>正文标题</w:t></w:r></w:p>
        <mc:AlternateContent>
          <mc:Choice Requires="w"><w:p><w:r><w:t>文本框内容</w:t></w:r></w:p></mc:Choice>
          <mc:Fallback><w:p><w:r><w:t>文本框内容</w:t></w:r></w:p></mc:Fallback>
        </mc:AlternateContent>
      </w:body>
    </w:document>""".encode("utf-8")
    payload = BytesIO()
    with ZipFile(payload, "w") as archive:
        archive.writestr("word/document.xml", xml)

    lines = _extract_docx_xml_lines(payload.getvalue())

    assert lines == ["正文标题", "文本框内容"]


def test_pdf_character_fragments_are_collapsed_without_touching_normal_lines() -> None:
    fragments = list("邮箱:test@example.com") + ["项目经历", "正常的完整项目描述"]

    collapsed = _collapse_fragmented_lines(fragments)

    assert collapsed[0] == "邮箱:test@example.com"
    assert collapsed[1:] == ["项目经历", "正常的完整项目描述"]


def test_rejects_scanned_or_empty_pdf() -> None:
    """Catches silent creation of an empty resume from an image-only PDF."""
    with pytest.raises(ResumeImportError, match="未提取到可用文字"):
        import_resume("scan.pdf", "application/pdf", make_blank_pdf())


def test_rejects_mismatched_extension_and_content_type() -> None:
    """Catches files disguised with a supported extension."""
    with pytest.raises(ResumeImportError, match="文件类型不匹配"):
        import_resume("resume.pdf", "text/plain", b"plain text")
