import json
from io import BytesIO

from docx import Document
from docx.shared import Cm, Pt

from resume_mvp.domain import (
    Fact,
    JobAnalysis,
    JobProject,
    LayoutProfile,
    ResumeDocument,
    ResumeVersion,
    SourcedText,
    utc_now,
)
from resume_mvp.exports import build_codex_handoff, build_docx, build_resume_json


def sample_resume() -> ResumeDocument:
    resume = ResumeDocument.blank()
    resume.basics.name = "张宁"
    resume.basics.email = "private@example.com"
    resume.basics.phone = "13800000000"
    resume.basics.target_role = SourcedText(value="后端工程师")
    resume.basics.summary = SourcedText(value="专注可靠 API")
    return resume


def sample_project() -> JobProject:
    now = utc_now()
    return JobProject(
        id="project-1",
        title="后端工程师",
        company_name="示例科技",
        application_type="experienced",
        job_description="负责 Python API",
        job_analysis=JobAnalysis(role_title="后端工程师"),
        selected_template_id="clear-single",
        created_at=now,
        updated_at=now,
    )


def test_docx_export_can_be_reopened_and_contains_current_resume() -> None:
    """Catches zero-byte or structurally invalid Word exports."""
    data = build_docx(sample_resume(), "clear-single")

    document = Document(BytesIO(data))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "张宁" in text
    assert "后端工程师" in text
    assert len(data) > 1_000


def test_json_export_excludes_secrets_and_contains_active_resume() -> None:
    """Catches provider credentials entering a portable resume export."""
    version = ResumeVersion(
        id="version-1",
        project_id="project-1",
        resume=sample_resume(),
        facts=[],
        reason="当前版本",
        created_at=utc_now(),
    )

    payload = json.loads(build_resume_json(sample_project(), version))

    assert payload["resume"]["basics"]["name"] == "张宁"
    assert payload["version_id"] == "version-1"
    assert "api_key" not in json.dumps(payload)


def test_codex_handoff_omits_contact_details_but_keeps_confirmed_facts() -> None:
    """Catches unnecessary contact data being sent in a Codex handoff."""
    fact = Fact(
        id="fact-1",
        category="工作经历",
        statement="使用 Python 开发订单 API",
        source_type="manual",
        user_confirmed=True,
    )

    markdown = build_codex_handoff(
        sample_project(),
        sample_project().job_analysis,
        sample_resume(),
        [fact],
    )

    assert "使用 Python 开发订单 API" in markdown
    assert "private@example.com" not in markdown
    assert "13800000000" not in markdown
    assert "不得虚构事实" in markdown


def test_campus_docx_uses_compact_layout_tokens() -> None:
    """Catches one-page DOCX exports rendering with social-hire spacing."""
    campus = Document(BytesIO(build_docx(sample_resume(), "clear-single", "campus")))
    experienced = Document(BytesIO(build_docx(sample_resume(), "clear-single", "experienced")))

    assert campus.sections[0].top_margin < experienced.sections[0].top_margin
    assert campus.styles["Normal"].font.size == Pt(9)
    assert campus.sections[0].top_margin <= Cm(1.3)


def test_docx_export_prefers_valid_imported_template_tokens() -> None:
    """Catches an uploaded DOCX template being discarded during export."""
    resume = sample_resume()
    resume.layout_profile = LayoutProfile(
        source_kind="docx",
        imported=True,
        font_family="FangSong",
        heading_font_family="FangSong",
        accent_color="#7A3E8E",
        base_font_size=11,
    )

    document = Document(BytesIO(build_docx(resume, "clear-single", "experienced")))

    assert document.styles["Normal"].font.name == "FangSong"
    assert document.styles["Normal"].font.size == Pt(11)
    assert str(document.paragraphs[0].runs[0].font.color.rgb) == "7A3E8E"
