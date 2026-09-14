from __future__ import annotations

import base64
from pathlib import Path

import fitz
import pytest
from fastapi.testclient import TestClient

from resume_mvp.main import create_app
from resume_mvp.preview import (
    PreviewConversionError,
    convert_docx_to_pdf,
    count_pdf_pages,
    render_pdf_page_pngs,
)
from resume_mvp.providers.e2e import E2EProvider
from tests.test_export_api import create_project_with_resume


def _client(tmp_path: Path) -> TestClient:
    return TestClient(create_app(data_dir=tmp_path, test_providers={"test": E2EProvider()}))


def _tiny_pdf(pages: int = 1) -> bytes:
    document = fitz.open()
    for _ in range(pages):
        document.new_page()
    payload = document.tobytes()
    document.close()
    return payload


def test_count_pdf_pages() -> None:
    assert count_pdf_pages(_tiny_pdf(3)) == 3


def test_preview_pdf_endpoint_returns_pdf_and_page_count(tmp_path: Path, monkeypatch) -> None:
    client = _client(tmp_path)
    project_id = create_project_with_resume(client)
    pdf = _tiny_pdf(2)

    monkeypatch.setattr(
        "resume_mvp.api.exports.convert_docx_to_pdf",
        lambda _docx: pdf,
    )
    monkeypatch.setattr(
        "resume_mvp.api.exports.count_pdf_pages",
        lambda _pdf: 2,
    )

    response = client.post(
        f"/api/projects/{project_id}/preview/pdf",
        json={"resume": None, "template_id": None},
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.headers["x-resume-page-count"] == "2"
    assert response.content[:4] == b"%PDF"


def test_preview_pdf_accepts_live_resume_body(tmp_path: Path, monkeypatch) -> None:
    client = _client(tmp_path)
    project_id = create_project_with_resume(client)
    version = client.get(f"/api/projects/{project_id}/versions").json()[0]
    resume = version["resume"]
    resume["basics"]["name"] = "预览草稿名"

    captured: dict[str, object] = {}

    def fake_convert(docx_bytes: bytes) -> bytes:
        captured["docx"] = docx_bytes
        return _tiny_pdf(1)

    monkeypatch.setattr("resume_mvp.api.exports.convert_docx_to_pdf", fake_convert)
    monkeypatch.setattr("resume_mvp.api.exports.count_pdf_pages", lambda _pdf: 1)

    response = client.post(
        f"/api/projects/{project_id}/preview/pdf",
        json={"resume": resume, "template_id": "clear-single"},
    )

    assert response.status_code == 200
    assert captured["docx"]
    from io import BytesIO
    from docx import Document
    text = "\n".join(p.text for p in Document(BytesIO(captured["docx"])).paragraphs)
    assert "预览草稿名" in text


def test_preview_pdf_unavailable_when_conversion_fails(tmp_path: Path, monkeypatch) -> None:
    client = _client(tmp_path)
    project_id = create_project_with_resume(client)

    def boom(_docx: bytes) -> bytes:
        raise PreviewConversionError("未找到 LibreOffice（soffice），无法生成 PDF 预览")

    monkeypatch.setattr("resume_mvp.api.exports.convert_docx_to_pdf", boom)

    response = client.post(f"/api/projects/{project_id}/preview/pdf", json={})
    assert response.status_code == 503
    assert response.json()["detail"]["code"] == "PREVIEW_UNAVAILABLE"


def test_preview_pages_endpoint_returns_png_base64(tmp_path: Path, monkeypatch) -> None:
    client = _client(tmp_path)
    project_id = create_project_with_resume(client)
    pdf = _tiny_pdf(2)
    png = b"\x89PNG\r\n\x1a\npage"

    monkeypatch.setattr("resume_mvp.api.exports.convert_docx_to_pdf", lambda _docx: pdf)
    monkeypatch.setattr("resume_mvp.api.exports.count_pdf_pages", lambda _pdf: 2)
    monkeypatch.setattr("resume_mvp.api.exports.render_pdf_page_pngs", lambda _pdf: [png, png])

    response = client.post(f"/api/projects/{project_id}/preview/pages", json={})

    assert response.status_code == 200
    body = response.json()
    assert body["page_count"] == 2
    assert len(body["pages"]) == 2
    assert body["pages"][0] == base64.b64encode(png).decode("ascii")


def test_render_pdf_page_pngs() -> None:
    pages = render_pdf_page_pngs(_tiny_pdf(2), dpi=72)
    assert len(pages) == 2
    assert pages[0][:8] == b"\x89PNG\r\n\x1a\n"


@pytest.mark.integration
def test_convert_docx_to_pdf_with_libreoffice() -> None:
    from io import BytesIO
    from docx import Document

    buffer = BytesIO()
    document = Document()
    document.add_heading("集成预览", 0)
    document.add_paragraph("用于验证 LibreOffice 转换链路。")
    document.save(buffer)

    pdf = convert_docx_to_pdf(buffer.getvalue())
    assert pdf[:4] == b"%PDF"
    assert count_pdf_pages(pdf) >= 1
