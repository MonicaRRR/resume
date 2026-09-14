from pathlib import Path
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from resume_mvp.main import create_app
from resume_mvp.providers.e2e import E2EProvider


def _client(tmp_path: Path) -> TestClient:
    return TestClient(create_app(data_dir=tmp_path, test_providers={"test": E2EProvider()}))


def _seed_profile(client: TestClient) -> None:
    client.put(
        "/api/profile",
        json={
            "resume": {
                "basics": {
                    "name": "张宁",
                    "email": "",
                    "phone": "",
                    "location": "",
                    "target_role": {"value": "后端工程师", "source_fact_ids": [], "origin": "manual", "confidence": 1},
                    "summary": {"value": "", "source_fact_ids": [], "origin": "manual", "confidence": 1},
                },
                "education": [],
                "work_experience": [{
                    "id": "w1",
                    "company": "示例科技",
                    "title": "实习生",
                    "start_date": "",
                    "end_date": "",
                    "bullets": [{"value": "使用 Python 开发 API", "source_fact_ids": [], "origin": "manual", "confidence": 1}],
                }],
                "projects": [],
                "skills": [{"id": "s1", "name": "技能", "items": [{"value": "Python", "source_fact_ids": [], "origin": "manual", "confidence": 1}]}],
                "certificates": [],
                "awards": [],
                "custom_sections": [],
                "section_order": ["basics", "work_experience", "projects", "education", "skills"],
                "layout_profile": {
                    "source_kind": "builtin", "font_family": "", "heading_font_family": "",
                    "accent_color": "", "base_font_size": None, "line_height": None, "columns": 1, "imported": False,
                },
            }
        },
    )


def create_project_with_resume(client: TestClient) -> str:
    _seed_profile(client)
    project = client.post(
        "/api/projects",
        json={
            "title": "后端工程师",
            "company_name": "示例科技",
            "application_type": "experienced",
            "job_description": "负责 Python API",
        },
    ).json()
    client.post(
        f"/api/projects/{project['id']}/resume/import",
        files={"file": ("resume.txt", "张宁\n后端工程师\n使用 Python 开发 API", "text/plain")},
    )
    return project["id"]


def test_export_routes_return_reopenable_files_and_handoff(tmp_path: Path) -> None:
    """Catches attachment routes returning the wrong version or media type."""
    client = _client(tmp_path)
    project_id = create_project_with_resume(client)

    docx_response = client.post(f"/api/projects/{project_id}/export/docx")
    assert docx_response.status_code == 200
    assert docx_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    path = tmp_path / "export.docx"
    path.write_bytes(docx_response.content)
    assert "张宁" in "\n".join(p.text for p in Document(path).paragraphs)

    json_response = client.get(f"/api/projects/{project_id}/export/json")
    assert json_response.status_code == 200
    assert json_response.json()["resume"]["basics"]["name"] == "张宁"

    handoff = client.post(f"/api/projects/{project_id}/codex-handoff")
    assert handoff.status_code == 200
    assert "不得虚构事实" in handoff.json()["markdown"]


def test_campus_docx_export_allowed_when_content_exceeds_one_page(tmp_path: Path) -> None:
    """Campus overflow is a soft warning; export must still succeed."""
    client = _client(tmp_path)
    _seed_profile(client)
    project = client.post(
        "/api/projects",
        json={
            "title": "校招后端工程师",
            "company_name": "",
            "application_type": "campus",
            "job_description": "负责 Python API",
        },
    ).json()
    long_text = "负责接口设计、稳定性治理与性能优化，" * 180
    client.post(
        f"/api/projects/{project['id']}/resume/import",
        files={"file": ("resume.txt", f"张宁\n后端工程师\n{long_text}", "text/plain")},
    )

    response = client.post(f"/api/projects/{project['id']}/export/docx")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def test_docx_export_uses_live_draft_basics(tmp_path: Path) -> None:
    """Catches Word export ignoring birthday/gender from the current editor draft."""
    client = _client(tmp_path)
    project_id = create_project_with_resume(client)
    version = client.get(f"/api/projects/{project_id}/versions").json()[0]
    resume = version["resume"]
    resume["basics"]["gender"] = "女"
    resume["basics"]["birthday"] = "2001-08-08"
    resume["basics"]["wechat"] = "draft-wechat"
    resume["basics"]["location"] = "杭州"

    response = client.post(
        f"/api/projects/{project_id}/export/docx",
        json={"resume": resume, "template_id": "clear-single"},
    )

    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + "\n".join(paragraph.text for paragraph in cell.paragraphs)
    assert "生日：2001-08-08" in text
    assert "性别：女" in text
    assert "微信：draft-wechat" in text
    assert "现居：杭州" in text
